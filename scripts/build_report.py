"""Build the submission report from real screenshots and recorded development steps."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'submission'
QA=OUT/'qa'
QA.mkdir(exist_ok=True)
doc=Document()
section=doc.sections[0]
section.top_margin=section.bottom_margin=Inches(.65)
section.left_margin=section.right_margin=Inches(.7)
section.page_width=Inches(8.5);section.page_height=Inches(11)
for name in ['Normal','Title','Heading 1','Heading 2','Caption']:
    style=doc.styles[name];style.font.name='Arial';style.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(10)
doc.styles['Normal'].paragraph_format.space_after=Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing=1.13
doc.styles['Title'].font.size=Pt(25)
doc.styles['Heading 1'].font.size=Pt(17)
doc.styles['Heading 2'].font.size=Pt(12)
doc.styles['Caption'].font.size=Pt(9)
doc.styles['Caption'].paragraph_format.space_after=Pt(12)
doc.core_properties.title='Account Takeover Investigation Project'
doc.core_properties.subject='Week 1 Path 3 project submission'
doc.core_properties.author=''
def p(text,style=None):return doc.add_paragraph(text,style)
def h(text):doc.add_heading(text,level=1)
def figure(name,box,caption,width=7.1):
    image=Image.open(OUT/'screenshots'/name)
    if box:image=image.crop(box)
    path=QA/(name.replace('.png','-report.png'));image.save(path)
    para=doc.add_paragraph();para.paragraph_format.space_after=Pt(4)
    para.add_run().add_picture(str(path),width=Inches(width))
    p(caption,'Caption')
def table(headers,rows):
    t=doc.add_table(rows=1,cols=len(headers));t.style='Light Shading Accent 1'
    for c,x in zip(t.rows[0].cells,headers):c.text=x
    for row in rows:
        for c,x in zip(t.add_row().cells,row):c.text=str(x)
    for cell in t.rows[0].cells:
        tcPr=cell._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EDF2F2');tcPr.append(shade)
        for r in cell.paragraphs[0].runs:r.font.bold=True;r.font.color.rgb=RGBColor(0,0,0)
    for row in t.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(4)
                for r in para.runs:r.font.size=Pt(9)
    return t

p('Account Takeover Investigation Project','Title')
p('Week 1 project  |  Path 3  |  ATO Investigator')
h('Project overview')
p('ATO Investigator is a functional data application that helps a fraud analyst connect login behavior, credential changes, and transfers into an account investigation. The application combines an overview, prioritized review queue, explainable event timeline, case notes, and a detection-rule playground.')
p('The problem is that a single unusual login can represent either a takeover or a legitimate customer change. The app brings related activity together and exposes the reasons for each alert, making the evidence and alternative explanations easier to inspect.')
h('Dataset and application')
p('The project uses 1,578 fictional events across 80 accounts from October 1 to 11, 2026 UTC. A Python standard-library generator uses seed 42 to reproduce the CSV and JSON. Events include successful and failed logins, password resets, beneficiary additions, and transfers. All amounts are USD; no real customer data is used.')
figure('01-overview.png',(0,0,1440,810),'Figure 1  The working overview shows account counts, review exposure, activity trends, and observed signals.')
p('The runtime is HTML, CSS, and JavaScript. It needs only a static HTTP server and has no application package dependencies. Synthetic data generation uses Python; the detection tests use the Node.js built-in test runner.')

doc.add_page_break()
h('Investigating an account takeover')
p('Opening ACC-0001 reveals a sequence on October 11: four failed logins, a password reset, a successful login from a new device and country, a beneficiary addition, and a $6,322 transfer. The analyst can switch between all activity and signal events, inspect evidence, record a review outcome, and export the timeline.')
figure('02-investigation.png',(234,290,1440,1090),'Figure 2  ACC-0001 exposes the login and transfer signals. All activity also includes the preceding failures and credential change.')
h('Explainable chronological scoring')
p('At a successful login, a new device and a new country add 20 points each. A recent failed-login burst adds 25; a recent password reset adds 15. A large transfer after a suspicious login on the same device adds 45. A large transfer to a newly added beneficiary adds 35. The account score is the maximum event score, capped at 100; it is not a fraud probability.')
p('Rules use only earlier observed events. The first successful login establishes the initial baseline. Labels and scenario descriptions are reserved for evaluation and the optional ground-truth reveal. Cases such as travel and password recovery intentionally create legitimate alerts.')
h('Investigator workflow')
p('The review queue supports account search, priority, signal reasons, exposure, and review status. Case notes are stored locally in the current browser and can be exported as JSON. CSV uploads are validated and analyzed locally. The application does not block transactions or contact customers.')

doc.add_page_break()
h('Evaluating detection tradeoffs')
p('The detection playground compares default and current rule settings on the same account cohort. Precision and recall are account-level metrics and exclude accounts without labels. Four deliberately low-signal takeovers remain undetected in both configurations shown below.')
table(['Metric','Default threshold 40','Threshold 70'],[['Accounts alerted',22,8],['Takeovers caught',8,8],['Legitimate accounts flagged',14,0],['Takeovers missed',4,4],['Precision','36.4%','100%'],['Recall','66.7%','66.7%']])
p('These results are designed examples on a synthetic dataset, not evidence of production fraud performance. Flagged transfer value measures rule-flagged transfers rather than confirmed fraud loss.')
figure('03-rules.png',(234,295,1440,1230),'Figure 3  Adjustable rule settings and the confusion matrix make the alert workload and detection limitations visible.',width=6.7)
p('Global dates select accounts with activity in the chosen UTC period; country selects the first-successful-login country. Scoring, timelines, and evaluation use full account history. The daily chart alone restricts events to selected dates, so rule context is preserved.')

doc.add_page_break()
h('Vibe coding workflow')
p('The project began with the supplied assignment guidelines and a request for fraud-domain ideas. After selecting account takeover, the build request was:')
p('“I want to pick account takeover project and would need help building the project end to end for submission.”')
for title,text in [
('1 Define the scope','Select a browser app with four views, synthetic login and transfer data, explainable rules, and no real customer information.'),
('2 Generate the dataset','Create reproducible routine behavior, strong and low-signal takeovers, legitimate travel, password recovery, and new-phone scenarios.'),
('3 Build the logic and interface','Separate detection and CSV validation from the UI. Add charts, a paginated review queue, account timelines, rules, notes, and exports.'),
('4 Verify and iterate','Run nine engine tests and thirteen browser checks. Review screenshots, improve mobile table scrolling and timeline controls, and update slider results on release for reliable dragging.'),
('5 Prepare submission assets','Capture the working views, document the actual build steps, prepare a live-demo narration script, and package the source code.')]:
    doc.add_heading(title,level=2);p(text)
h('Validation and development observations')
p('All nine engine tests pass. They verify expected outcomes, threshold effects, independence from labels, exclusion of future activity, expired time windows, missing labels, CSV quoting, invalid uploads, and spreadsheet-formula escaping in exports. Browser checks verify navigation, filters, timelines, note persistence, rule changes, CSV upload and export, empty states, invalid dates, and mobile layout.')
p('The local server initially required socket-binding permission. The browser automation initially could not find its default browser executable; using the installed Chrome executable resolved that setup issue. These were environment issues rather than changes to the fraud logic.')
h('Learnings and next steps')
p('The implementation demonstrates why labels must be kept separate from detection inputs, why metrics need explicit definitions, and why a useful demo includes false positives and missed attacks. Future improvements could add MFA and recovery events, session grouping, analyst assignment, and evaluation on approved historical data.')
p('For final submission, add screenshots of the actual Codex conversation and your own observations. Import this report into Google Docs, verify the public GitHub code link, record the live walkthrough in five minutes or less, and submit those three links through the assignment form. Suggested follow-up prompts and the narration script are separate from the actual development history.')

target=OUT/'Account_Takeover_Project_Report.docx'
doc.save(target)
print(target)
