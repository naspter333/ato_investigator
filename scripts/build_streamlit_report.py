"""Create the revised report from actual prompts, artifacts, and live screenshots."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from PIL import Image
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'submission';QA=OUT/'qa';QA.mkdir(exist_ok=True)
J=json.loads((OUT/'process_journal.json').read_text())
D=Document();sec=D.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.7)
for name in ['Normal','Title','Heading 1','Heading 2','Caption']:
    sty=D.styles[name];sty.font.name='Arial';sty.font.color.rgb=RGBColor(0,0,0)
D.styles['Normal'].font.size=Pt(10)
D.styles['Normal'].paragraph_format.line_spacing=1.08
D.styles['Normal'].paragraph_format.space_after=Pt(6)
D.styles['Title'].font.size=Pt(25)
D.styles['Heading 1'].font.size=Pt(17)
D.styles['Heading 2'].font.size=Pt(11)
D.styles['Heading 2'].paragraph_format.space_before=Pt(8)
D.styles['Caption'].font.size=Pt(8)
D.styles['Caption'].paragraph_format.space_after=Pt(8)
D.core_properties.title='Account Takeover Streamlit Project'
D.core_properties.author=''
def p(text,style=None):return D.add_paragraph(text,style)
def h(text):D.add_heading(text,1)
def fig(name,caption,crop=None,width=4.5):
    im=Image.open(OUT/'screenshots'/'streamlit'/name)
    if crop:im=im.crop(crop)
    dest=QA/('report-'+name);im.save(dest)
    para=D.add_paragraph();para.paragraph_format.space_after=Pt(4)
    para.add_run().add_picture(str(dest),width=Inches(width))
    p(caption,'Caption')
def step(x):
    D.add_heading(f'{x["number"]} {x["title"]}',2);p(x['detail'])
def page():D.add_page_break()

p('Account Takeover Streamlit Project','Title')
p('Week 1  |  Path 3  |  ATO Investigator')
h('Problem and final application')
p('Fraud analysts need to connect login attempts, device and country changes, credential recovery, beneficiary additions, and transfers into an account story. A suspicious login can represent a takeover or a legitimate customer change. ATO Investigator makes the evidence and alternative explanations visible in a working Streamlit data application.')
p('The final app uses Python, Streamlit, pandas, and Plotly. Its five workspace views cover the overview, account investigation, detection playground, data methodology, and vibe coding build journal. The Python entrypoint is app.py. The earlier static HTML demo is retained only as development history.')
h('Dataset and reproducibility')
p('A Python standard-library generator with seed 42 creates 1,578 fictional events across 80 accounts for October 1–11, 2026 UTC. Scenarios include 8 strong-signal takeovers, 4 low-signal takeovers, 8 legitimate travelers, 6 password recoveries, 6 new phones, and 48 routine accounts. Amounts are USD; there is no real customer data.')
fig('01-overview.jpg','Figure 1  The live Streamlit overview shows 80 accounts, 22 alerts, $35,118 of flagged transfer exposure, and 66.7 percent recall.',crop=(0,135,682,740),width=4.4)
p('Run locally with streamlit run app.py after installing requirements.txt. Streamlit runs Python on a server and presents its interface in a browser; this replaces the original static frontend.')

page();h('Actual prompts used during vibe coding')
p('The following are the actual human messages used to direct this project. They preserve the original wording. Implementation decisions and commands are documented separately in the build steps, rather than presented as invented user prompts.')
for i,item in enumerate(J['prompts'],1):
    D.add_heading(f'Prompt {i} {item["stage"]}',2);p(item['text'])
fig('04-process-prompts.jpg','Figure 2  Screenshot of the live Streamlit journal rendering the original project prompts. This is a journal capture, not an original Codex chat UI screenshot.',crop=(0,345,682,685),width=6.6)
p('Process evidence includes actual messages, current source files, and saved executed command outputs. Automated capture of the original Codex interface was blocked. If original chat UI images are required by the evaluator, add those manually; the journal images are explicitly identified here.')

page();h('Full build breakdown from idea to prototype')
for x in J['steps'][:6]:step(x)
fig('06-process-steps.jpg','Figure 3  The running build journal records scope, dataset generation, and the original prototype stages.',crop=(0,345,682,710),width=6.4)
p('The initial report emphasized finished-app screenshots. The user correction led to the Streamlit conversion and this expanded process-focused report. The history is retained so the iteration is visible.')

page();h('Full build breakdown for the Streamlit version')
for x in J['steps'][6:]:step(x)
fig('07-process-code.jpg','Figure 4  The live code evidence view displays fraud_engine.py, including rule defaults and input validation.',crop=(0,345,682,740),width=6.2)
p('Source artifacts: app.py, fraud_engine.py, requirements.txt, .streamlit/config.toml, tests/test_fraud_engine.py, and tests/test_streamlit_app.py. The build journal reads these real files rather than displaying mock code.')

page();h('Testing debugging and validation evidence')
p('The final Python validation run passes 14 tests. Nine engine tests verify expected outcomes, threshold effects, label independence, exclusion of future activity, expired time windows, absent labels, CSV quoting, invalid uploads, and export formula protection. Five Streamlit AppTest workflows cover overview filters and empty results, investigations and notes, rule comparisons, data methodology, and the process journal.')
fig('08-process-tests.jpg','Figure 5  Screenshot of the executed command output displayed in the live build journal. The saved full output is in submission/streamlit_validation.json.',crop=(0,75,682,635),width=5.2)
h('Observed bug and correction')
p('Adding a second radio selector exposed a test-harness error: a positional widget index targeted the wrong radio group. The failing output is retained in streamlit_validation_failure.json. Stable widget keys workspace_view and journal_evidence corrected the selection; the final run passes all 14 tests. This is a recorded iteration, not an invented debugging example.')
p('The local environment also required native network permission for dependencies and socket permission for the Streamlit server. The installed dependency versions are pinned. AppTest warnings about missing ScriptRunContext are expected in bare test mode; the final test exit code is zero.')

page();h('Working investigation and detection tradeoffs')
p('In ACC-0001, failed logins are followed by a password reset, a successful login on a new device in a new country, a new beneficiary, and a $6,322 transfer. The account score is 80. The analyst can inspect all activity or signal events, reveal optional synthetic ground truth, record an outcome, and export the timeline.')
fig('02-investigation.jpg','Figure 6  The running Streamlit account investigation exposes score, transfer values, and chronological activity.',crop=(0,140,682,700),width=4.4)
p('Default threshold 40 produces 8 takeovers caught, 14 legitimate accounts flagged, and 4 takeovers missed: precision 36.4 percent and recall 66.7 percent. Threshold 70 removes the 14 benign alerts in this sample, retaining the same 8 caught takeovers and 4 missed takeovers. These engineered synthetic results are not a production performance claim.')
h('Learnings and interpretation')
p('Labels and scenario text are excluded from scoring. Rules use earlier events only. Dates select an account cohort while preserving historical rule context. Flagged transfer value is review exposure rather than confirmed loss. Benign travel and password recovery expose false positives; low-signal takeovers show that changing thresholds alone cannot create missing evidence.')
h('Demo and submission')
p('The demo script covers the problem, overview, ACC-0001, benign ACC-0013, threshold comparison, data schema, and build journal within five minutes. Submit a Google Doc imported from this report, a narrated live video, and the public repository link: https://github.com/naspter333/ato_investigator.')
p('Streamlit CSV uploads are processed on the Python server. Local execution processes them on the user computer; hosted execution processes them on the host. Case notes last for the current session and should be exported before ending it. The narrated recording, Google Doc import, and final form submission remain separate completion steps.')
D.save(OUT/'Account_Takeover_Project_Report.docx')
print('Revised Streamlit report saved')
